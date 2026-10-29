"""Generate the presentation from the storyboard and image assets.

PowerPoint generation needs only python-pptx. Optional PNG previews use the
bundled presentation renderer supplied by ChatGPT or Codex.
"""

import argparse
from copy import deepcopy
import json
import os
from pathlib import Path
import re
import subprocess

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK,
)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SLIDES_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = SLIDES_DIR.parent.parent
SLIDE_IDS = [str(number) for number in range(1, 12)] + [f"A{number}" for number in range(1, 23)]
FONT = "Arial"
BACKGROUND = "FFFFFF"
INK = "163846"
TEAL = "167F83"
SECONDARY = "4B626B"
PALE = "EAF4F3"
RULE = "D8E5E5"


def pixels(value):
    """Convert the original generator's 96-DPI coordinates to slide units."""
    return Inches(value / 96)


def font_size(value):
    return Pt(value * 72 / 96)


def style_font(font, size, *, color=INK, bold=False, underline=False):
    font.name = FONT
    font.size = font_size(size)
    font.color.rgb = RGBColor.from_string(color)
    font.bold = bold
    font.underline = underline


def text(slide, name, value, x, y, width, height, size, *, color=INK,
         bold=False, alignment=PP_ALIGN.LEFT, underline=False):
    shape = slide.shapes.add_textbox(
        pixels(x), pixels(y), pixels(width), pixels(height)
    )
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraphs = value.split("\n") if isinstance(value, str) else [value]
    for index, content in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = alignment
        paragraph.space_before = paragraph.space_after = Pt(0)
        style_font(paragraph.font, size, color=color, bold=bold, underline=underline)
        runs = [content] if isinstance(content, str) else content
        for item in runs:
            run = paragraph.add_run()
            run.text = item if isinstance(item, str) else item[0]
            style_font(run.font, size, color=color,
                       bold=bold if isinstance(item, str) else item[1],
                       underline=underline)
    return shape


def notes_for(storyboard, slide_id):
    start = re.search(rf"^### {re.escape(slide_id)} ", storyboard, re.MULTILINE)
    if start is None:
        raise ValueError(f"Storyboard section missing for {slide_id}.")
    section = storyboard[start.start():]
    section = re.split(r"\n### ", section, maxsplit=1)[0]
    notes = re.search(
        r"\*\*Speaker notes[^\n]*\*\*\n\n([\s\S]*?)\n\n\*\*(?:Transition|Close)",
        section,
    )
    if notes is None:
        raise ValueError(f"Speaker notes missing for {slide_id}.")
    paragraphs = re.split(r"\n\s*\n", notes[1].replace("“", "").replace("”", ""))
    return "\n\n".join(" ".join(paragraph.split()) for paragraph in paragraphs)


def new_slide(presentation, storyboard, slide_id, title=None, *, two_lines=False):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BACKGROUND)
    if title is not None:
        text(slide, f"{slide_id} title", title, 72, 48, 1136,
             112 if two_lines else 66, 44 if slide_id.isdigit() else 46, bold=True)
    if slide_id != "1":
        text(slide, "Slide ID", slide_id, 1166, 675, 42, 22, 16,
             color=SECONDARY, alignment=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = notes_for(storyboard, slide_id)
    return slide


def footnote(slide, value):
    text(slide, "Scope", value, 72, 674, 1070, 24, 18, color=SECONDARY)


def rect(slide, name, x, y, width, height, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, pixels(x), pixels(y), pixels(width), pixels(height)
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    return shape


def picture(slide, path, alt, x, y, width, height):
    with Image.open(path) as image:
        image_width, image_height = image.size
    scale = min(width / image_width, height / image_height)
    fitted_width, fitted_height = image_width * scale, image_height * scale
    shape = slide.shapes.add_picture(
        str(path), pixels(x + (width - fitted_width) / 2),
        pixels(y + (height - fitted_height) / 2),
        width=pixels(fitted_width), height=pixels(fitted_height),
    )
    shape.name = path.stem
    shape._element.nvPicPr.cNvPr.set("descr", alt)
    return shape


def cell_borders(cell):
    # python-pptx does not expose table-cell borders through its public API.
    properties = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        line = OxmlElement(f"a:{edge}")
        line.set("w", str(font_size(1)))
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:srgbClr")
        color.set("val", RULE)
        fill.append(color)
        line.append(fill)
        dash = OxmlElement("a:prstDash")
        dash.set("val", "solid")
        line.append(dash)
        properties.append(line)


def native_table(slide, values, y, widths, row_height=66, *, result_column=False):
    table = slide.shapes.add_table(
        len(values), len(widths), pixels(72), pixels(y), pixels(sum(widths)),
        pixels(48 + row_height * (len(values) - 1)),
    ).table
    table.first_row = False
    table.horz_banding = False
    for column, width in zip(table.columns, widths):
        column.width = pixels(width)
    for row_index, row_values in enumerate(values):
        table.rows[row_index].height = pixels(48 if row_index == 0 else row_height)
        for column_index, value in enumerate(row_values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                PALE if row_index == 0 else BACKGROUND
            )
            cell.margin_left = cell.margin_right = pixels(18)
            cell.margin_top = cell.margin_bottom = pixels(10)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.auto_size = MSO_AUTO_SIZE.NONE
            cell.text_frame.word_wrap = True
            is_result = result_column and row_index > 0 and column_index == 1
            size = 23 if row_index == 0 else 32 if is_result else 26
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.space_before = paragraph.space_after = Pt(0)
            style_font(paragraph.font, size, bold=row_index == 0 or is_result,
                       color=TEAL if is_result else INK)
            for run in paragraph.runs:
                style_font(run.font, size, bold=row_index == 0 or is_result,
                           color=TEAL if is_result else INK)
            cell_borders(cell)
    return table


def survey_bullets(slide):
    shape = text(slide, "Survey and project data", "", 72, 148, 744, 500, 30)
    items = [
        ("MEPS:", " Medical Expenditure Panel Survey"),
        ("Data:", " 2023 Household Component (HC-251)"),
        ("Sample:", " 14,768 representing ~260 million U.S. adults"),
        ("Target:", " Annual out-of-pocket healthcare costs"),
        ("Features:", " 26 inputs across demographics, socioeconomics, "
         "health profile, chronic conditions, and limitations"),
        ("Survey weights:", " Used in EDA, model training, and evaluation "
         "to reflect the population"),
    ]
    for index, (label, value) in enumerate(items):
        paragraph = (shape.text_frame.paragraphs[0] if index == 0
                     else shape.text_frame.add_paragraph())
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(18)
        style_font(paragraph.font, 30)
        # Native bullets keep wrapped lines aligned with the text.
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Pt(18)))
        properties.set("indent", str(-Pt(12)))
        bullet_font = OxmlElement("a:buFont")
        bullet_font.set("typeface", FONT)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        properties.append(bullet_font)
        properties.append(bullet)
        for content, bold in ((label, True), (value, False)):
            run = paragraph.add_run()
            run.text = content
            style_font(run.font, 30, bold=bold)


def native_chart(slide, categories, values, name, x, y, width, height, *,
                 horizontal=False, color=INK, gap_width=140):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(name, values, number_format="0.0%")
    chart_type = (XL_CHART_TYPE.BAR_CLUSTERED if horizontal
                  else XL_CHART_TYPE.COLUMN_CLUSTERED)
    shape = slide.shapes.add_chart(
        chart_type, pixels(x), pixels(y), pixels(width), pixels(height), chart_data
    )
    shape.name = name
    chart = shape.chart
    # Normalize python-pptx's signed axis IDs for renderers that require unsigned IDs.
    for axis_id in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"):
        axis_id.set("val", str(int(axis_id.get("val")) % (2 ** 32)))
    style_font(chart.font, 24)
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = gap_width
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.show_value = True
    plot.data_labels.show_category_name = False
    plot.data_labels.show_series_name = False
    plot.data_labels.show_legend_key = False
    plot.data_labels.number_format = "0.0%"
    plot.data_labels.number_format_is_linked = False
    style_font(plot.data_labels.font, 27 if horizontal else 28, bold=True)
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    series.format.line.fill.background()
    category_axis, value_axis = chart.category_axis, chart.value_axis
    style_font(category_axis.tick_labels.font, 25 if horizontal else 24)
    style_font(value_axis.tick_labels.font, 22 if horizontal else 20)
    category_axis.has_major_gridlines = False
    value_axis.minimum_scale = 0
    value_axis.maximum_scale = 0.6 if horizontal else 1
    value_axis.major_unit = 0.2 if horizontal else 0.25
    value_axis.tick_labels.number_format = "0%"
    value_axis.tick_labels.number_format_is_linked = False
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(RULE)
    value_axis.major_gridlines.format.line.width = font_size(1)
    for axis in (category_axis, value_axis):
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.color.rgb = RGBColor.from_string(RULE)
        axis.format.line.width = font_size(1)
    if horizontal:
        # Keep the first category at the top, matching the original slide.
        category_axis.reverse_order = True
        series.points[0].format.fill.solid()
        series.points[0].format.fill.fore_color.rgb = RGBColor.from_string("A7B6BC")
    return chart


def build_presentation():
    storyboard = (SLIDES_DIR / "storyboard.md").read_text(encoding="utf-8")
    presentation = Presentation()
    presentation.slide_width = pixels(1280)
    presentation.slide_height = pixels(720)

    # Match hyperlink theme colors to the appendix links.
    theme_part = presentation.slide_master.part.part_related_by(RT.THEME)
    theme = etree.fromstring(theme_part.blob)
    color_scheme = theme.find(".//" + qn("a:clrScheme"))
    for color_name in ("a:hlink", "a:folHlink"):
        color_scheme.find(qn(color_name) + "/" + qn("a:srgbClr")).set("val", SECONDARY)
    theme_part._blob = etree.tostring(
        theme, xml_declaration=True, encoding="UTF-8", standalone=True,
    )

    cover = new_slide(presentation, storyboard, "1")
    picture(cover, WORKSPACE_DIR / "assets/header.png",
            "Medical Cost Planner project header.", 72, 72, 1136, 422)
    text(cover, "Subtitle", "Predicting Out-of-Pocket Healthcare Costs with Machine Learning",
         72, 522, 1136, 56, 34, bold=True)
    text(cover, "Presenter", "Jens Bender", 72, 594, 400, 36, 28)
    text(cover, "Event details", "[Event / setting] · [Date]", 72, 632, 1000, 32, 24,
         color=SECONDARY)

    problem = new_slide(presentation, storyboard, "2",
                        "How much should I set aside for healthcare?")
    text(problem, "Challenge label", "The challenge", 72, 520, 556, 36, 30, bold=True)
    text(problem, "Budgeting challenge", "Difficult to plan next year's out-of-pocket\n"
         "costs and HSA/FSA contributions.", 72, 560, 556, 72, 30)
    text(problem, "Aim label", "The aim", 652, 520, 556, 36, 30, bold=True)
    text(problem, "Intended estimate", "A useful ballpark estimate from questions "
         "people can answer from memory.", 652, 560, 556, 72, 30)
    picture(problem, SLIDES_DIR / "assets/budget-planning.png",
            "Illustration of an adult considering a budget with a planner and calculator.",
            72, 120, 1136, 379)
    text(problem, "Problem statement appendix label", "Appendix:",
         72, 674, 90, 24, 18, color=SECONDARY)
    healthcare_appendix_link = text(
        problem, "Healthcare costs appendix link", "U.S. healthcare cost explainer",
        168, 674, 704, 24, 18, color=SECONDARY, underline=True,
    )

    data = new_slide(presentation, storyboard, "3",
                     "MEPS links accessible inputs to observed spending")
    survey_bullets(data)
    picture(data, SLIDES_DIR / "assets/household-survey.png",
            "Illustration of a household survey interview: a respondent shows an empty "
            "wallet beside bills, while an interviewer listens with a laptop.",
            868, 148, 340, 453.333)
    text(data, "Data appendix label", "Appendix:",
         72, 674, 90, 24, 18, color=SECONDARY)
    appendix_link = text(data, "MEPS appendix link", "MEPS overview",
                         168, 674, 135, 24, 18, color=SECONDARY, underline=True)
    text(data, "Data appendix separator", "|", 314, 674, 12, 24, 18,
         color=SECONDARY, alignment=PP_ALIGN.CENTER)
    features_appendix_link = text(data, "Input features appendix link", "Features",
                                  340, 674, 90, 24, 18,
                                  color=SECONDARY, underline=True)

    distribution = new_slide(presentation, storyboard, "4")
    text(distribution, "4 title", "Out-of-pocket costs: 20% of adults account for ~80%",
         72, 48, 1136, 66, 44, bold=True)
    lorenz_path = WORKSPACE_DIR / "figures/eda/lorenz_curve.png"
    with Image.open(lorenz_path) as lorenz_image:
        image_width, image_height = lorenz_image.size
    # Crop the embedded title while preserving the axes and annotations.
    title_crop = 70
    chart_height = 580
    chart_width = image_width / (image_height - title_crop) * chart_height
    lorenz = distribution.shapes.add_picture(
        str(lorenz_path), pixels((1280 - chart_width) / 2 + 64), pixels(120),
        pixels(chart_width), pixels(chart_height),
    )
    lorenz.crop_top = title_crop / image_height
    lorenz.name = "Survey-weighted Lorenz curve"
    lorenz._element._nvXxPr.cNvPr.set(
        "descr", "Lorenz curve of annual out-of-pocket costs: 22.3% of adults "
        "have zero costs; the highest-spending 20% account for 79.3% of costs "
        "and the highest-spending 1% for 20.6%.",
    )
    metrics = text(distribution, "Evaluation metrics",
                   "Primary metric\nMdAE\nSupporting\nMAE",
                   75.8, 126.9, 184, 170, 24)
    for paragraph, heading, spacing in zip(
        metrics.text_frame.paragraphs, (True, False, True, False), (4, 18, 4, 0),
    ):
        paragraph.space_after = font_size(spacing)
        for run in paragraph.runs:
            style_font(run.font, 24, bold=heading)
    text(distribution, "EDA appendix label", "Appendix:", 72, 588, 164, 22, 18,
         color=SECONDARY)
    outlier_appendix_link = text(
        distribution, "Outlier appendix link", "Outlier analysis",
        72, 614, 164, 24, 18, color=SECONDARY, underline=True,
    )
    distributions_appendix_link = text(
        distribution, "Distributions appendix link", "Distributions",
        72, 640, 164, 24, 18, color=SECONDARY, underline=True,
    )
    relationships_appendix_link = text(
        distribution, "Feature-target relationships appendix link",
        "Feature–target relations",
        72, 666, 240, 24, 18, color=SECONDARY, underline=True,
    )

    preprocessing = new_slide(presentation, storyboard, "5")
    preprocessing.background.fill.fore_color.rgb = RGBColor.from_string("020617")
    for shape in preprocessing.shapes:
        if shape.name == "Slide ID":
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor.from_string("94A3B8")
    picture(preprocessing, SLIDES_DIR / "assets/preprocessing-pipeline.png",
            "Preprocessing pipeline fitted on training data: categorical input "
            "standardization, missing-value validation and imputation, health-burden "
            "feature derivation, scaling, and encoding.", 0, 0, 1280, 720)
    text(preprocessing, "Preprocessing appendix label", "Appendix:",
         20, 648, 200, 22, 18, color="94A3B8")
    preparation_appendix_link = text(
        preprocessing, "Data preparation appendix link", "Data preparation",
        20, 674, 200, 24, 18, color="94A3B8", underline=True,
    )

    selection = new_slide(presentation, storyboard, "6",
                          "Model selection: median error was not enough")
    model_table = native_table(selection, [
        ["Tuned model", "MdAE", "Overfitting", "MAE"],
        ["Elastic Net", "$159", "+7.9%", "$1,051"],
        ["Random Forest", "$228", "+3.8%", "$964"],
        ["XGBoost", "$242", "+6.2%", "$954"],
    ], 142, [420, 216, 250, 250])
    for row, column in ((1, 1), (2, 2), (3, 3)):
        for run in model_table.cell(row, column).text_frame.paragraphs[0].runs:
            run.font.bold = True
    text(selection, "Validation note", "Note: Survey-weighted validation metrics. "
         "Overfitting is the percentage change in MdAE from training to validation.",
         72, 402, 1136, 24, 18, color=SECONDARY)
    selection_bullets = text(selection, "Model selection rationale", "",
                             72, 456, 1136, 180, 26)
    for index, (label, value) in enumerate([
        ("Elastic Net:", " Lowest MdAE, but never predicts over ~$217 and increasingly "
         "underpredicts."),
        ("Choose XGBoost:", " Predicts up to ~$2,100 with less systematic underprediction."),
        ("Switch to quantile regression:", " Errors generally widen at higher predicted "
         "costs, so add budgeting ranges to communicate that uncertainty."),
    ]):
        paragraph = (selection_bullets.text_frame.paragraphs[0] if index == 0
                     else selection_bullets.text_frame.add_paragraph())
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(10)
        style_font(paragraph.font, 26)
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Pt(18)))
        properties.set("indent", str(-Pt(12)))
        bullet_font = OxmlElement("a:buFont")
        bullet_font.set("typeface", FONT)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        properties.append(bullet_font)
        properties.append(bullet)
        for content, bold in ((label, True), (value, False)):
            run = paragraph.add_run()
            run.text = content
            style_font(run.font, 26, bold=bold)
    text(selection, "Model selection appendix label", "Appendix:",
         72, 674, 90, 24, 18, color=SECONDARY)
    baseline_appendix_link = text(
        selection, "Baseline models appendix link", "Baseline models",
        168, 674, 140, 24, 18, color=SECONDARY, underline=True,
    )
    text(selection, "Model selection appendix separator", "|", 314, 674, 12, 24, 18,
         color=SECONDARY, alignment=PP_ALIGN.CENTER)
    tuning_appendix_link = text(
        selection, "Hyperparameter tuning appendix link", "Hyperparameter tuning",
        342, 674, 210, 24, 18, color=SECONDARY, underline=True,
    )
    text(selection, "Model diagnostics appendix separator", "|",
         540, 674, 12, 24, 18, color=SECONDARY, alignment=PP_ALIGN.CENTER)
    residuals_appendix_link = text(
        selection, "Heteroscedasticity appendix link", "Heteroscedasticity",
        568, 674, 190, 24, 18, color=SECONDARY, underline=True,
    )
    text(selection, "Subgroup reliability appendix separator", "|",
         738, 674, 12, 24, 18, color=SECONDARY, alignment=PP_ALIGN.CENTER)
    reliability_appendix_link = text(
        selection, "Subgroup reliability appendix link", "Subgroup reliability",
        766, 674, 205, 24, 18, color=SECONDARY, underline=True,
    )

    quantiles = new_slide(presentation, storyboard, "7",
                          "Quantile regression: budgeting for uncertainty")
    text(quantiles, "Plan-around label", "Plan-around estimate (q50)",
         166, 204, 500, 42, 30,
         bold=True, alignment=PP_ALIGN.CENTER)
    text(quantiles, "Safety label", "Safety cushion (q90)", 848, 204, 360, 42, 30,
         bold=True, alignment=PP_ALIGN.CENTER)
    # This axis is conceptual; positions do not encode a person's dollar values.
    rect(quantiles, "Cost axis", 100, 304, 1080, 2, INK)
    rect(quantiles, "Typical range q25 to q75", 250, 282, 440, 44, PALE)
    for x, label in ((250, "q25"), (416, "q50"), (690, "q75"), (1028, "q90")):
        rect(quantiles, f"Quantile marker {label}", x, 274, 3, 60, INK)
        text(quantiles, label, label, x - 45, 353, 93, 42, 30, alignment=PP_ALIGN.CENTER)
    brace = quantiles.shapes.add_shape(
        MSO_SHAPE.RIGHT_BRACE, pixels(459), pixels(200), pixels(22), pixels(440)
    )
    brace.name = "Typical range brace"
    brace.rotation = 90
    brace.fill.background()
    brace.line.color.rgb = RGBColor.from_string(SECONDARY)
    brace.line.width = Pt(1.2)
    brace._element.spPr.append(OxmlElement("a:effectLst"))
    text(quantiles, "Range label", "Typical range (q25–q75)", 200, 450, 540, 42, 30,
         bold=True, alignment=PP_ALIGN.CENTER)
    text(quantiles, "Coverage target", "Targets 50% coverage", 250, 496, 440, 40, 28,
         color=SECONDARY, alignment=PP_ALIGN.CENTER)
    text(quantiles, "Upper coverage target", "Targets 90% coverage", 848, 496, 360, 40, 28,
         color=SECONDARY, alignment=PP_ALIGN.CENTER)

    results = new_slide(presentation, storyboard, "8")
    text(results, "8 title", "XGBoost quantile audit: all release gates passed",
         72, 48, 1136, 66, 44, bold=True)
    results_table = native_table(results, [
        ["Test metric", "Estimate (95% CI)", "Release gate"],
        ["Plan-around estimate (q50): MdAE", "$240 [$215, $279]", "< $500"],
        ["Typical range (q25–q75): coverage", "47.3% [44.0%, 50.6%]", "45%–55%"],
        ["Safety cushion (q90): coverage", "91.0% [89.2%, 92.6%]", "85%–95%"],
        ["Typical range (q25–q75): mean width", "$912 [$875, $955]", "< $1,500"],
        ["Safety cushion (q50–q90): mean width", "$2,032 [$1,964, $2,108]", "< $3,500"],
    ], 136, [568, 336, 232], row_height=58, result_column=True)
    for row in list(results_table.rows)[1:]:
        paragraph = row.cells[1].text_frame.paragraphs[0]
        estimate, interval = paragraph.text.split(" [", 1)
        paragraph.clear()
        run = paragraph.add_run()
        run.text = estimate
        style_font(run.font, 26, bold=True)
        run = paragraph.add_run()
        run.text = f" [{interval}"
        style_font(run.font, 26)
    text(results, "Test metrics note", "Note: Survey-weighted test metrics. "
         "95% percentile bootstrap confidence intervals from 1,000 resamples of "
         "test-set rows.", 72, 486, 1136, 24, 18, color=SECONDARY)
    audit_summary_items = [
        ("Value over population and age-group baselines:",
         " better typical range and safety cushion."),
        ("Subgroup reliability and fairness:",
         " add planning notes for groups at risk of underbudgeting."),
    ]
    for index, (label, detail) in enumerate(audit_summary_items):
        bullet_shape = text(results, f"Audit summary {index + 1}",
                            [(label, True), detail], 72, 542 + index * 58,
                            1136, 36, 26)
        properties = bullet_shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
        properties.set("marL", str(Pt(18)))
        properties.set("indent", str(-Pt(12)))
        bullet_font = OxmlElement("a:buFont")
        bullet_font.set("typeface", FONT)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        properties.append(bullet_font)
        properties.append(bullet)
    text(results, "Audit appendix label", "Appendix:",
         72, 674, 88, 24, 18, color=SECONDARY)
    audit_appendix_links = text(
        results, "Audit appendix links",
        ["Value over simple baselines", "  |  ", "Subgroup reliability",
         "  |  ", "Subgroup fairness"],
        160, 674, 960, 24, 18, color=SECONDARY,
    )
    (comparison_appendix_link, _, final_reliability_appendix_link,
     _, final_fairness_appendix_link) = (
        audit_appendix_links.text_frame.paragraphs[0].runs
    )
    comparison_appendix_link.font.underline = True
    final_reliability_appendix_link.font.underline = True
    final_fairness_appendix_link.font.underline = True

    example = new_slide(presentation, storyboard, "9", "Example prediction output")
    expanded = new_slide(presentation, storyboard, "10", "Example prediction output")
    for slide, slide_id in ((example, "9"), (expanded, "10")):
        next(shape for shape in slide.shapes if shape.name == f"{slide_id} title").width = pixels(568)
        text(slide, "Example profile", "(high cost profile: 69 years, arthritis)", 642, 66, 578, 36, 26, color=SECONDARY)
        example_border = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, pixels(48), pixels(118), pixels(1184), pixels(520),
        )
        example_border.name = "Prediction output border"
        example_border.fill.background()
        example_border.line.color.rgb = RGBColor.from_string(RULE)
        example_border.line.width = Pt(1)
        example_border._element.spPr.append(OxmlElement("a:effectLst"))
    text(example, "Output heading", "Your estimated out-of-pocket costs for next year",
         72, 132, 1136, 38, 28, bold=True)
    for slide, label_y, value_y, value_size in (
        (example, 190, 228, 42),
        (expanded, 150, 186, 34),
    ):
        for label, value, x, width in (
            ("Plan around", "$895", 72, 320),
            ("Typical range", "$400–$2,025", 464, 352),
            ("Safety cushion", "$3,852", 888, 320),
        ):
            text(slide, f"{label} label", label, x, label_y, width, 32, 24, bold=True)
            text(slide, f"{label} value", value, x, value_y, width, 54, value_size,
                 color=TEAL, bold=True)
    text(example, "Output interpretation", "Plan around gives you an estimate of your "
         "typical costs. The typical range shows how much those costs usually vary. "
         "The safety cushion helps you budget for a higher-cost year. Costs can still be higher.",
         72, 294, 1136, 76, 22)
    text(example, "Planning note heading", "Planning note",
         72, 378, 1136, 32, 24, bold=True)
    text(example, "Planning note", "Costs for profiles like yours can vary a lot. "
         "To reduce the risk of underbudgeting, consider planning closer to the "
         "safety cushion.", 72, 414, 1136, 56, 22)
    for index, label in enumerate((
        "Which answers shaped your estimate?", "How you compare to others",
    )):
        expand_link = text(example, f"Expand details {index + 1}",
                           [(label, True), " (click to expand)"],
                           72, 510 + index * 54, 1136, 32, 24)
        label_run, hint_run = expand_link.text_frame.paragraphs[0].runs
        label_run.font.underline = True
        hint_run.font.color.rgb = RGBColor.from_string(SECONDARY)
        expand_link.click_action.target_slide = expanded
    text(expanded, "Comparison heading", "How you compare to others",
         720, 272, 488, 36, 26, bold=True)
    comparison_rows = [
        ("Your plan-around estimate", "$895"),
        ("Typical American", "$268"),
        ("Typical for ages 65+", "$657"),
    ]
    comparison_table = native_table(expanded, comparison_rows, 316, [350, 138], row_height=46)
    expanded.shapes[-1].name = "Cost comparison table"
    expanded.shapes[-1].left = pixels(720)
    comparison_table.rows[0].height = pixels(46)
    for row_index, row_values in enumerate(comparison_rows):
        for column_index, value in enumerate(row_values):
            cell = comparison_table.cell(row_index, column_index)
            cell.fill.fore_color.rgb = RGBColor.from_string(BACKGROUND)
            cell.margin_left = cell.margin_right = 0
            cell.margin_top = cell.margin_bottom = pixels(4)
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                line = cell._tc.get_or_add_tcPr().find(qn(f"a:{edge}"))
                for child in list(line):
                    line.remove(child)
                line.append(OxmlElement("a:noFill"))
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT if column_index else PP_ALIGN.LEFT
            style_font(paragraph.font, 26)
            for run in paragraph.runs:
                style_font(run.font, 26)

    text(expanded, "SHAP heading", "Which answers shaped your estimate?",
         72, 272, 536, 36, 26, bold=True)
    shap_rows = [
        ("Arthritis: Yes", "↑ +$124"),
        ("Joint pain: Yes", "↑ +$110"),
        ("High blood pressure: Yes", "↑ +$94"),
        ("Age: 69", "↑ +$92"),
        ("Insurance: Any private", "↑ +$80"),
    ]
    shap_table = native_table(expanded, shap_rows, 316, [350, 138], row_height=52)
    expanded.shapes[-1].name = "SHAP contributions table"
    shap_table.rows[0].height = pixels(52)
    for row_index, row_values in enumerate(shap_rows):
        for column_index, value in enumerate(row_values):
            cell = shap_table.cell(row_index, column_index)
            cell.fill.fore_color.rgb = RGBColor.from_string(BACKGROUND)
            cell.margin_left = cell.margin_right = 0
            cell.margin_top = cell.margin_bottom = pixels(4)
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                line = cell._tc.get_or_add_tcPr().find(qn(f"a:{edge}"))
                for child in list(line):
                    line.remove(child)
                line.append(OxmlElement("a:noFill"))
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.clear()
            paragraph.alignment = PP_ALIGN.RIGHT if column_index else PP_ALIGN.LEFT
            size = 26
            if column_index == 0:
                feature, answer = value.split(":", 1)
                for content, is_bold in ((feature + ":", True), (answer, False)):
                    run = paragraph.add_run()
                    run.text = content
                    style_font(run.font, size, bold=is_bold)
            else:
                run = paragraph.add_run()
                run.text = value
                style_font(run.font, size)

    text(expanded, "SHAP appendix label", "Appendix:",
         72, 674, 86, 24, 18, color=SECONDARY)
    shap_details_link = text(expanded, "SHAP details appendix link", "SHAP details",
                             160, 674, 160, 24, 18, color=SECONDARY, underline=True)

    next_steps = new_slide(presentation, storyboard, "11", "Model evaluation is complete;\n"
                           "app development comes next", two_lines=True)
    text(next_steps, "Implemented", "Implemented", 72, 190, 700, 38, 30, bold=True)
    text(next_steps, "Planned", "Planned", 895, 190, 313, 38, 30, bold=True)
    text(next_steps, "Training heading", "Training artifacts", 72, 265, 330, 38, 29, bold=True)
    text(next_steps, "Training tools", "DVC stages\nMLflow experiments\nEvaluation artifacts",
         72, 310, 330, 127, 27)
    text(next_steps, "Training to inference", "→", 418, 295, 64, 58, 42)
    text(next_steps, "Shared heading", "Shared inference", 510, 265, 320, 38, 29, bold=True)
    text(next_steps, "Inference modules", "Prediction and SHAP\nReusable modules\nUnit tests",
         510, 310, 320, 127, 27)
    text(next_steps, "Inference to app", "→", 822, 295, 64, 58, 42)
    text(next_steps, "App heading", "FastAPI / Gradio", 895, 265, 313, 38, 29, bold=True)
    text(next_steps, "Remaining app work", "Integration and latency\nUser evaluation\n"
         "Aggregate monitoring", 895, 310, 313, 127, 27)
    text(next_steps, "Validation heading", "Next validation priority", 72, 478, 1136, 38, 28,
         bold=True)
    text(next_steps, "Validation work", "Check feature timing and performance on a later "
         "survey year", 72, 520, 1136, 38, 28)
    text(next_steps, "Closing lesson", "Evaluate the estimates needed for the user decision",
         72, 605, 1136, 44, 32, bold=True)

    healthcare_overview = new_slide(presentation, storyboard, "A1")
    picture(healthcare_overview, WORKSPACE_DIR / "assets/infographic_healthcare_costs.png",
            "Overview of U.S. healthcare payments, out-of-pocket costs, insurance "
            "terms, HSA/FSA planning, and the project target.", 20, 0, 1240, 660)
    healthcare_return_link = text(
        healthcare_overview, "Return to problem statement", "Back",
        72, 674, 800, 24, 18, color=SECONDARY, underline=True,
    )
    healthcare_appendix_link.click_action.target_slide = healthcare_overview
    healthcare_return_link.click_action.target_slide = problem

    meps_overview = new_slide(presentation, storyboard, "A2")
    picture(meps_overview, WORKSPACE_DIR / "assets/infographic_meps_data.jpg",
            "MEPS household, provider, and employer survey components and the 2023 "
            "data used in this project.", 20, 0, 1240, 660)
    return_link = text(meps_overview, "Return to data slide", "Back",
                       72, 674, 800, 24, 18, color=SECONDARY, underline=True)
    appendix_link.click_action.target_slide = meps_overview
    return_link.click_action.target_slide = data

    feature_values = new_slide(presentation, storyboard, "A3")
    text(feature_values, "A3 title", "Preprocessor input features",
         72, 48, 1136, 66, 44, bold=True)
    feature_values_table = native_table(feature_values, [
        ["Feature", "Type", "Possible values"],
        ["Age", "Numerical", "18\u201385 years"],
        ["Sex", "Binary", "Male, female"],
        ["Region", "Categorical", "Northeast, Midwest, South, West"],
        ["Marital status", "Categorical", "Married, widowed, divorced, separated, never married"],
        ["Family income", "Numerical", "1 = Poor/Negative, 2 = Near Poor, 3 = Low, 4 = Middle, 5 = High"],
        ["Family size", "Numerical", "1\u201314"],
        ["Education", "Categorical", "None, GED, high school, bachelor's, master's, doctorate, other"],
        ["Employment status", "Binary", "Employed, not employed"],
        ["Insurance", "Categorical", "Any private, public only, uninsured"],
        ["Physical health", "Numerical", "1 = Excellent, 2 = Very good, 3 = Good, 4 = Fair, 5 = Poor"],
        ["Mental health", "Numerical", "1 = Excellent, 2 = Very good, 3 = Good, 4 = Fair, 5 = Poor"],
    ], 142, [270, 160, 706], row_height=42)

    feature_answers = new_slide(presentation, storyboard, "A4")
    text(feature_answers, "A4 title", "Preprocessor input features: yes/no",
         72, 48, 1136, 66, 44, bold=True)
    feature_answers_table = native_table(feature_answers, [
        ["Feature group", "Features"],
        ["Chronic conditions", "High blood pressure, high cholesterol, diabetes,\n"
         "coronary heart disease, stroke, cancer, arthritis, asthma"],
        ["Limitations", "ADL help: personal care, such as bathing or dressing\n"
         "IADL help: daily tasks, such as managing bills or medication\n"
         "Walking limitation, cognitive limitation, joint pain"],
        ["Other", "Smoker\n"
         "Usual source of care: regular doctor or clinic\n"
         "Recent life transition (derived): marital or employment changes"],
    ], 142, [280, 856], row_height=120)
    for row, height in zip(feature_answers_table.rows, [48, 120, 156, 156]):
        row.height = pixels(height)
    for feature_table in (feature_values_table, feature_answers_table):
        for row_index, row in enumerate(feature_table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.margin_left = cell.margin_right = pixels(12)
                cell.margin_top = cell.margin_bottom = pixels(6)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.space_before = paragraph.space_after = Pt(0)
                    is_group = feature_table is feature_answers_table and column_index == 0
                    style_font(paragraph.font, 24, bold=row_index == 0 or is_group)
                    for run in paragraph.runs:
                        style_font(run.font, 24, bold=row_index == 0 or is_group)
    for feature_slide in (feature_values, feature_answers):
        feature_return_link = text(
            feature_slide, "Return to data slide", "Back",
            72, 674, 800, 24, 18, color=SECONDARY, underline=True,
        )
        feature_return_link.click_action.target_slide = data
    features_appendix_link.click_action.target_slide = feature_values

    outliers = new_slide(
        presentation, storyboard, "A5",
        "Outliers: greater health burden and higher cost risk",
    )
    text(outliers, "Outlier detection",
         [("Detection: ", True), ("Isolation Forest (flagging 5%)", False)],
         72, 142, 1136, 42, 30)
    profile_table = native_table(outliers, [
        ["Profiling", "Inliers", "Outliers"],
        ["In top 1% of spenders", "0.9%", "3.5%"],
        ["Age (median)", "47", "66"],
        ["Chronic conditions (median)", "1", "4"],
        ["Functional limitations (median)", "0", "3"],
        ["Private insurance", "68%", "32%"],
        ["Walking limitations", "9%", "75%"],
        ["Cognitive limitations", "3%", "61%"],
    ], 206, [680, 228, 228], row_height=42)
    for row_index in range(len(profile_table.rows)):
        for column_index in range(len(profile_table.columns)):
            cell = profile_table.cell(row_index, column_index)
            cell.margin_top = cell.margin_bottom = pixels(6)
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if column_index == 0 else PP_ALIGN.CENTER
            style_font(paragraph.font, 26, bold=row_index == 0)
            for run in paragraph.runs:
                style_font(run.font, 26, bold=row_index == 0)
    text(outliers, "Outlier treatment",
         [("Treatment: ", True),
          ("Retained all outliers as plausible, valuable training examples.", False)],
         72, 583.56, 1136, 42, 30)
    outlier_return_link = text(
        outliers, "Return to cost distribution", "Back",
        72, 674, 800, 24, 18, color=SECONDARY, underline=True,
    )
    outlier_appendix_link.click_action.target_slide = outliers
    outlier_return_link.click_action.target_slide = distribution
    for slide_id, filename, description in [
        ("A6", "numerical_distributions.png", "Population distributions of age, "
         "family size, physical health, and mental health."),
        ("A7", "categorical_distributions.png", "Population distributions of "
         "region, marital status, employment, insurance, education, and family income."),
        ("A8", "binary_distributions.png", "Population distributions of binary "
         "features, including chronic conditions, limitations, and access to care."),
    ]:
        feature_distributions = new_slide(presentation, storyboard, slide_id)
        picture(feature_distributions, WORKSPACE_DIR / "figures/eda" / filename,
                description, 0, 0, 1280, 720)
        distribution_return_link = text(
            feature_distributions, "Return to cost distribution", "Back",
            20, 674, 100, 24, 18, color=SECONDARY, underline=True,
        )
        distribution_return_link.click_action.target_slide = distribution
        if slide_id == "A6":
            distributions_appendix_link.click_action.target_slide = feature_distributions
    for slide_id, filename, description in [
        ("A9", "correlation_heatmap.png", "Population Spearman rank correlations "
         "between out-of-pocket costs and numerical, ordinal, and binary features."),
        ("A10", "numerical_feature_target_relationships.png", "Relationships between "
         "out-of-pocket costs and age, family size, physical health, and mental health."),
        ("A11", "categorical_feature_target_relationships.png", "Out-of-pocket cost "
         "distributions by region, marital status, employment, insurance, education, "
         "and family income."),
        ("A12", "binary_feature_target_relationships.png", "Out-of-pocket cost "
         "distributions by binary features, including chronic conditions, limitations, "
         "and access to care."),
    ]:
        feature_relationships = new_slide(presentation, storyboard, slide_id)
        picture(feature_relationships, WORKSPACE_DIR / "figures/eda" / filename,
                description, 0, 0, 1280, 720)
        relationships_return_link = text(
            feature_relationships, "Return to cost distribution", "Back",
            20, 674, 100, 24, 18, color=SECONDARY, underline=True,
        )
        relationships_return_link.click_action.target_slide = distribution
        if slide_id == "A9":
            relationships_appendix_link.click_action.target_slide = feature_relationships
    preparation = new_slide(presentation, storyboard, "A13")
    text(preparation, "A13 title", "Data preparation",
         72, 48, 1136, 66, 44, bold=True)
    text(preparation, "Preparation heading", "Prepare MEPS data",
         72, 148, 1136, 42, 30, bold=True)
    text(preparation, "Preparation steps", "Select variables and adults, "
         "standardize missing values, and group sparse categories.",
         72, 195, 1136, 42, 28)
    text(preparation, "Split heading", "Stratified split by out-of-pocket costs",
         72, 283.31, 1136, 42, 30, bold=True)
    text(preparation, "Split rationale", "Keep zero-cost and rare high-cost cases "
         "represented across splits.", 72, 329.31, 1136, 42, 28)
    text(preparation, "Zero-cost group label", "Zero costs",
         72, 391.33, 140, 36, 25, alignment=PP_ALIGN.CENTER)
    text(preparation, "Lower-cost groups", "Lower costs",
         218, 391.33, 210, 36, 25)
    text(preparation, "Extreme-cost groups", "Extreme costs",
         995, 391.33, 213, 36, 25, alignment=PP_ALIGN.RIGHT)
    rect(preparation, "Separate zero-cost group", 72, 435.33, 140, 54, INK)
    group_x = 218
    for index, width in enumerate((290, 240, 180, 126, 80, 44)):
        rect(preparation, f"Positive-cost percentile group {index + 1}",
             group_x, 435.33, width, 54, INK)
        group_x += width + 6
    text(preparation, "Tail grouping", "Finer percentile groups toward the expensive tail",
         218, 503.32, 990, 36, 25)
    text(preparation, "Split allocation",
         "Random split within each group: 80% train, 10% validation, 10% test",
         72, 554.31, 1136, 38, 25)
    preparation_return_link = text(
        preparation, "Return to preprocessing pipeline", "Back",
        72, 674, 800, 24, 18, color=SECONDARY, underline=True,
    )
    preparation_appendix_link.click_action.target_slide = preparation
    preparation_return_link.click_action.target_slide = preprocessing

    baselines = new_slide(presentation, storyboard, "A14")
    text(baselines, "A14 title", "Baseline models",
         72, 48, 1136, 66, 44, bold=True)
    baseline_table = native_table(baselines, [
        ["Model", "MdAE", "Overfitting", "MAE"],
        ["Elastic Net", "$163", "+6.6%", "$1,044"],
        ["Linear Regression", "$219", "+4.8%", "$998"],
        ["Random Forest", "$232", "+9.6%", "$958"],
        ["Median (benchmark)", "$248", "0.0%", "$1,041"],
        ["Decision Tree", "$271", "+1.5%", "$971"],
        ["XGBoost", "$281", "+98.0%", "$961"],
        ["Support Vector Machine", "$291", "+190.7%", "$1,027"],
        ["LLM (benchmark)", "$518", "N/A", "$1,168"],
    ], 142, [420, 216, 250, 250], row_height=48)
    for row in baseline_table.rows:
        for cell in row.cells:
            cell.margin_top = cell.margin_bottom = pixels(6)
    for row in (1, 3, 6):
        for run in baseline_table.cell(row, 0).text_frame.paragraphs[0].runs:
            run.font.bold = True
    for row, column in ((1, 1), (3, 3), (5, 2)):
        for run in baseline_table.cell(row, column).text_frame.paragraphs[0].runs:
            run.font.bold = True
    for row in (4, 8):
        for cell in baseline_table.rows[row].cells:
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.italic = True
    text(baselines, "Validation note", "Note: Survey-weighted validation metrics. "
         "Overfitting is the percentage change in MdAE from training to validation.",
         72, 602, 1136, 24, 18, color=SECONDARY)
    baseline_return_link = text(
        baselines, "Return to model selection", "Back",
        72, 674, 800, 24, 18, color=SECONDARY, underline=True,
    )
    baseline_appendix_link.click_action.target_slide = baselines
    baseline_return_link.click_action.target_slide = selection

    tuning = new_slide(presentation, storyboard, "A15")
    text(tuning, "A15 title", "Hyperparameter tuning",
         72, 48, 1136, 66, 44, bold=True)
    text(tuning, "Search and selection", "50 configurations per model; "
         "select the lowest weighted validation MdAE", 72, 142, 1136, 42, 28)
    tuning_columns = [
        ("Elastic Net", 72, 300, [
            "Squared terms: on/off",
            "Regularization: 0.01–1.0",
            "L1/L2 mix: 0–1",
        ]),
        ("Random Forest", 428, 348, [
            "Trees: 200, 300, 400",
            "Depth: 8–24",
            "Min. split size: 20–149",
            "Min. leaf size: 10–79",
            "Features: sqrt, log2, 30/50/70%",
            "Row sampling: 60%–100%",
        ]),
        ("XGBoost", 860, 348, [
            "Trees: 400, 600, 800",
            "Depth: 3–9",
            "Learning rate: 0.01–0.2",
            "Min. child weight: 1–19",
            "Row sampling: 60%–100%",
            "Feature sampling: 50%–100%",
            "L1 penalty: 0–5",
            "L2 penalty: 0–5",
        ]),
    ]
    for model, x, width, parameters in tuning_columns:
        text(tuning, f"{model} heading", model, x, 222, width, 42, 32, bold=True)
        for index, parameter in enumerate(parameters):
            text(tuning, f"{model} parameter {index + 1}", parameter,
                 x, 280 + index * 44, width, 34, 24)
    tuning_return_link = text(
        tuning, "Return to model selection", "Back",
        72, 674, 800, 24, 18, color=SECONDARY, underline=True,
    )
    tuning_appendix_link.click_action.target_slide = tuning
    tuning_return_link.click_action.target_slide = selection
    residuals = new_slide(presentation, storyboard, "A16")
    picture(residuals, WORKSPACE_DIR / "figures/evaluation/"
            "tuned_models_validation_heteroscedasticity.png",
            "Validation residuals versus predicted costs for tuned Elastic Net, "
            "Random Forest, and XGBoost, with survey-weighted median residuals "
            "and interquartile ranges.", 0, 0, 1280, 720)
    residuals_return_link = text(
        residuals, "Return to model selection", "Back",
        20, 674, 100, 24, 18, color=SECONDARY, underline=True,
    )
    residuals_appendix_link.click_action.target_slide = residuals
    residuals_return_link.click_action.target_slide = selection
    reliability = new_slide(presentation, storyboard, "A17")
    reliability_plot = picture(reliability, SLIDES_DIR / "assets/"
            "tuned_models_validation_subgroup_reliability.png",
            "Survey-weighted validation median absolute errors for tuned Elastic "
            "Net, Random Forest, and XGBoost across cost tiers, physical health, "
            "insurance types, and chronic condition counts.", 0, 0, 1280, 720)
    reliability_plot.top = pixels(32)
    reliability_return_link = text(
        reliability, "Return to model selection", "Back",
        20, 674, 100, 24, 18, color=SECONDARY, underline=True,
    )
    reliability_appendix_link.click_action.target_slide = reliability
    reliability_return_link.click_action.target_slide = selection

    comparison = new_slide(presentation, storyboard, "A18")
    text(comparison, "A18 title", "XGBoost quantile: value over simple baselines",
         72, 48, 1136, 66, 44, bold=True)
    text(comparison, "Baseline comparison introduction", "Compared XGBoost with "
         "population and age-group baselines based on training quantiles.\n"
         "The clearest gains were in the typical range and safety cushion.",
         72, 116, 1136, 70, 26)
    comparison_table = native_table(comparison, [
        ["Metric", "Versus population\nbaseline (95% CI)",
         "Versus age-group\nbaseline (95% CI)"],
        ["Plan-around estimate (q50):\nMdAE reduction", "3.4% [-12.7%, 13.1%]",
         "21.5% [10.1%, 30.1%]"],
        ["Plan-around estimate (q50):\nquantile skill score", "9.8% [7.8%, 11.9%]",
         "7.6% [5.9%, 9.6%]"],
        ["Typical range (q25–q75):\ninterval skill score", "11.2% [8.6%, 13.8%]",
         "9.0% [6.9%, 11.3%]"],
        ["Safety cushion (q90):\nquantile skill score", "15.6% [10.8%, 21.1%]",
         "14.3% [9.7%, 19.0%]"],
    ], 202, [536, 300, 300], row_height=76)
    comparison_table.rows[0].height = pixels(68)
    for row_index, row in enumerate(comparison_table.rows):
        for cell in row.cells:
            cell.margin_top = cell.margin_bottom = pixels(6)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_before = paragraph.space_after = Pt(0)
                style_font(paragraph.font, 23 if row_index == 0 else 26,
                           bold=row_index == 0)
                for run in paragraph.runs:
                    style_font(run.font, 23 if row_index == 0 else 26,
                               bold=row_index == 0)
    for row in list(comparison_table.rows)[1:]:
        for cell in list(row.cells)[1:]:
            paragraph = cell.text_frame.paragraphs[0]
            estimate, interval = paragraph.text.split(" [", 1)
            paragraph.clear()
            run = paragraph.add_run()
            run.text = estimate
            style_font(run.font, 26, bold=True)
            run = paragraph.add_run()
            run.text = f" [{interval}"
            style_font(run.font, 26)
    text(comparison, "Comparison note", "Note: Survey-weighted test metrics. "
         "Skill scores measure the percentage reduction in quantile or interval "
         "loss compared with each baseline.\n"
         "95% percentile bootstrap confidence intervals from 1,000 paired "
         "resamples of test-set rows.",
         72, 594, 1136, 48, 18, color=SECONDARY)
    comparison_return_link = text(
        comparison, "Return to final model audit", "Back",
        72, 674, 100, 24, 18, color=SECONDARY, underline=True,
    )
    comparison_return_link.click_action.target_slide = results

    final_reliability = new_slide(presentation, storyboard, "A19")
    final_reliability_plot = picture(final_reliability, SLIDES_DIR / "assets/"
            "xgb_quantile_test_subgroup_reliability.png",
            "Survey-weighted test median absolute error, coverage, and mean width "
            "for XGBoost quantile regression across actual-cost and predicted-Q90 tiers.",
            0, 0, 1280, 660)
    final_reliability_plot.top = pixels(32)
    final_reliability_return_link = text(
        final_reliability, "Return to final model audit", "Back",
        20, 674, 100, 24, 18, color=SECONDARY, underline=True,
    )
    final_fairness = new_slide(presentation, storyboard, "A20")
    final_fairness_plot = picture(final_fairness, SLIDES_DIR / "assets/"
            "xgb_quantile_test_subgroup_fairness.png",
            "Survey-weighted test median absolute error, coverage, and mean width "
            "for XGBoost quantile regression across mental health, family income, "
            "and education groups.", 32, 0, 1216, 704)
    final_fairness_plot.top = pixels(16)
    final_fairness_return_link = text(
        final_fairness, "Return to final model audit", "Back",
        20, 684, 100, 18, 14, color=SECONDARY, underline=True,
    )
    fairness_slide_id = next(shape for shape in final_fairness.shapes if shape.name == "Slide ID")
    fairness_slide_id.top = pixels(684)
    fairness_slide_id.height = pixels(18)
    for run in fairness_slide_id.text_frame.paragraphs[0].runs:
        style_font(run.font, 14, color=SECONDARY)
    for run, target in (
        (comparison_appendix_link, comparison),
        (final_reliability_appendix_link, final_reliability),
        (final_fairness_appendix_link, final_fairness),
    ):
        hyperlink = OxmlElement("a:hlinkClick")
        hyperlink.set(qn("r:id"), results.part.relate_to(target.part, RT.SLIDE))
        hyperlink.set("action", "ppaction://hlinksldjump")
        run._r.get_or_add_rPr().append(hyperlink)
    final_reliability_return_link.click_action.target_slide = results
    final_fairness_return_link.click_action.target_slide = results

    shap_details = new_slide(presentation, storyboard, "A21",
                             "SHAP explains how answers shape predictions")
    for heading, description, y, height in (
        ("Background data", "A reference sample drawn from training data. Its average "
         "plan-around estimate provides the starting point.", 148, 112),
        ("Feature contributions", "Show how each of the 27 answers moves the "
         "plan-around estimate up or down.", 264, 112),
        ("Why permutation SHAP", "Explains the full prediction pipeline using "
         "answers before preprocessing.", 380, 112),
        ("Limitations", "Predicted rather than actual costs,\n"
         "no causal interpretation, correlated features,\n"
         "background data dependence.", 496, 112),
    ):
        lines = description.split("\n")
        section = text(shap_details, heading, [(f"{heading}:", True), f" {lines[0]}"],
                       72, y, 680, height, 28)
        paragraph = section.text_frame.paragraphs[0]
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(pixels(24)))
        properties.set("indent", str(-pixels(16)))
        bullet_font = OxmlElement("a:buFont")
        bullet_font.set("typeface", FONT)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        properties.append(bullet_font)
        properties.append(bullet)
        for line in lines[1:]:
            paragraph.add_line_break()
            run = paragraph.add_run()
            run.text = line
            style_font(run.font, 28)
    shap_example_border = shap_details.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, pixels(784), pixels(148), pixels(448), pixels(276),
    )
    shap_example_border.name = "SHAP example border"
    shap_example_border.fill.background()
    shap_example_border.line.color.rgb = RGBColor.from_string(RULE)
    shap_example_border.line.width = Pt(1)
    shap_example_border._element.spPr.append(OxmlElement("a:effectLst"))
    shap_example = deepcopy(next(shape for shape in expanded.shapes
                                if shape.name == "SHAP contributions table")._element)
    shap_details.shapes._spTree.insert_element_before(shap_example, "p:extLst")
    shap_example_shape = shap_details.shapes[-1]
    shap_example_shape.left = pixels(808)
    shap_example_shape.top = pixels(156)
    shap_example_shape.width = pixels(400)
    shap_example_shape.table.columns[0].width = pixels(292)
    shap_example_shape.table.columns[1].width = pixels(108)
    for row in shap_example_shape.table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                style_font(paragraph.font, 24)
                for run in paragraph.runs:
                    style_font(run.font, 24, bold=bool(run.font.bold))
    shap_details_return_link = text(shap_details, "Return to example prediction", "Back",
                                    72, 674, 100, 24, 18, color=SECONDARY, underline=True)
    shap_details_link.click_action.target_slide = shap_details
    shap_details_return_link.click_action.target_slide = expanded

    shap_benchmark = new_slide(
        presentation, storyboard, "A22",
        "SHAP benchmarking: fastest configuration passing quality checks",
    )
    benchmark_title = next(shape for shape in shap_benchmark.shapes
                           if shape.name == "A22 title")
    benchmark_title.width = pixels(1160)
    for paragraph in benchmark_title.text_frame.paragraphs:
        for run in paragraph.runs:
            style_font(run.font, 36, bold=True)
    for heading, description, y, height, level in (
        ("Configurations", "Tested 50–300 background data rows and 1–12 permutation rounds "
         "against a larger reference configuration (500 rows, 24 rounds).", 148, 112, 0),
        ("Winner", "225 background rows and 1 permutation round.", 260, 40, 0),
        ("Quality checks", "Passed.", 312, 40, 0),
        ("Background validation", "8.8% difference from full training data (limit: 10%).",
         356, 76, 1),
        ("Explanation stability", "Minimum 4 of 5 top features matched the reference "
         "configuration. Median absolute contribution difference was $6.47 (limit: $25).",
         436, 112, 1),
        ("SHAP explanation latency", "P95 was 0.20 seconds.", 552, 40, 0),
    ):
        section = text(shap_benchmark, heading, [(f"{heading}:", True), f" {description}"],
                       72 + 32 * level, y, 804 - 32 * level, height, 28)
        paragraph = section.text_frame.paragraphs[0]
        paragraph.level = level
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(pixels(24)))
        properties.set("indent", str(-pixels(16)))
        bullet_font = OxmlElement("a:buFont")
        bullet_font.set("typeface", FONT)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "◦" if level else "•")
        properties.append(bullet_font)
        properties.append(bullet)
    picture(shap_benchmark, SLIDES_DIR / "assets/shap-benchmarking.png",
            "Illustrated homage to Lloyd Shapley adjusting two knobs while holding "
            "a stopwatch, with a PASS card on the test bench.", 928, 148, 280, 373.333)
    shap_benchmark_link = text(
        shap_details, "SHAP benchmarking appendix link", "SHAP benchmarking",
        190, 674, 300, 24, 18, color=SECONDARY, underline=True,
    )
    shap_benchmark_link.click_action.target_slide = shap_benchmark
    shap_benchmark_return_link = text(
        shap_benchmark, "Return to SHAP explanation details", "Back",
        72, 674, 100, 24, 18, color=SECONDARY, underline=True,
    )
    shap_benchmark_return_link.click_action.target_slide = shap_details
    return presentation


def render_previews(pptx_path):
    """Use the bundled renderer when the caller requests PNG previews."""
    required = ("SKILL_DIR", "RUNTIME_NODE", "RUNTIME_NODE_MODULES")
    missing = [name for name in required if not os.environ.get(name)]
    if missing:
        raise ValueError(f"PNG previews require: {', '.join(missing)}.")
    renderer = Path(os.environ["SKILL_DIR"]) / "container_tools/render_presentation.mjs"
    result = subprocess.run(
        [os.environ["RUNTIME_NODE"], str(renderer), "--input", str(pptx_path),
         "--output_dir", str(pptx_path.parent), "--scale", "1.5"],
        check=False, capture_output=True, text=True, encoding="utf-8",
    )
    if result.returncode:
        raise RuntimeError(f"PNG rendering failed:\n{result.stderr or result.stdout}")
    paths = json.loads(result.stdout)["paths"]
    if len(paths) != len(SLIDE_IDS):
        raise ValueError(f"Expected {len(SLIDE_IDS)} previews, got {len(paths)}.")
    for slide_id, path in zip(SLIDE_IDS, paths):
        Path(path).rename(pptx_path.parent / f"{slide_id}.png")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("revision", help="New revision name, for example v4")
    parser.add_argument("--render", action="store_true",
                        help="Also create PNG previews using the bundled renderer")
    args = parser.parse_args()
    if re.fullmatch(r"[a-zA-Z0-9_-]+", args.revision) is None:
        parser.error("Revision names can contain only letters, numbers, underscores and hyphens.")
    output_dir = SLIDES_DIR / "exports" / args.revision
    if output_dir.exists():
        parser.error(f"Revision already exists: {output_dir}. Choose a new name.")
    if args.render:
        missing = [name for name in ("SKILL_DIR", "RUNTIME_NODE", "RUNTIME_NODE_MODULES")
                   if not os.environ.get(name)]
        if missing:
            parser.error(f"PNG previews require: {', '.join(missing)}.")
    presentation = build_presentation()
    output_dir.mkdir(parents=True)
    pptx_path = output_dir / f"medical-cost-planner-{args.revision}.pptx"
    presentation.save(pptx_path)
    print(f"PowerPoint: {pptx_path}")
    if args.render:
        render_previews(pptx_path)
        print(f"PNG previews: {output_dir}")


if __name__ == "__main__":
    main()
